"""
End-to-end tests for pipeline orchestration, CLI, data enrichment, and Excel reporting.
"""
from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook
import pandas as pd
import pytest

from data.fetch import enrich_dataframe
from data.synthetic import gbm_series
from main import main
from pipeline.runner import run_pipeline
from pipeline.state import PipelineState
from reporting.excel_report import write_excel_report


class TestEndToEndPipeline:
    def test_full_pipeline_synthetic_and_excel_report(self, tmp_path: Path):
        """Run full pipeline offline on synthetic data and verify attempts log & Excel sheets."""
        state = PipelineState(
            tickers=["SYN_TREND_A", "SYN_TREND_B", "SYN_NOISE"],
            risk_tolerance="moderate",
            max_attempts=10,
            confidence_threshold=0.55,
            dsr_threshold=0.70,
            seed=42,
            synthetic=True,
            n_mc_paths=80,
        )
        run_pipeline(state)

        assert len(state.attempts_log) > 0
        for att in state.attempts_log:
            if not att.approved:
                assert att.rejection_stage is not None and att.rejection_stage != ""
                assert att.rejection_reason is not None and att.rejection_reason != ""

        report_path = tmp_path / "test_report.xlsx"
        write_excel_report(state, report_path)
        assert report_path.exists()

        wb = load_workbook(report_path)
        assert wb.sheetnames == ["Summary", "Risk Detail", "Portfolio"]
        # Summary should have header + 1 row per attempt
        assert wb["Summary"].max_row == len(state.attempts_log) + 1

    def test_pure_noise_universe_yields_zero_approved(self):
        """A pure-noise synthetic universe must result in zero approved strategies."""
        state = PipelineState(
            tickers=["NOISE_1", "NOISE_2", "NOISE_3"],
            risk_tolerance="conservative",
            max_attempts=10,
            confidence_threshold=0.70,
            dsr_threshold=0.90,
            seed=101,
            synthetic=True,
            pure_noise=True,
            n_mc_paths=80,
        )
        run_pipeline(state)
        assert len(state.attempts_log) > 0
        assert len(state.approved_strategies) == 0
        assert state.portfolio_status == "insufficient strategies"
        for att in state.attempts_log:
            assert not att.approved
            assert att.rejection_stage in {"backtest", "validation", "risk", "scoring"}
            assert bool(att.rejection_reason)

    def test_skips_short_history_tickers(self):
        """Tickers with history shorter than MIN_HISTORY_BARS are skipped in data step."""
        short_df = enrich_dataframe(gbm_series(n=100, seed=1))
        long_df = enrich_dataframe(gbm_series(n=600, mu=0.25, sigma=0.12, seed=2))
        state = PipelineState(
            market_data={"SHORT_TICKER": short_df, "LONG_TICKER": long_df},
            max_attempts=5,
            synthetic=True,
            n_mc_paths=50,
        )
        run_pipeline(state)
        assert "SHORT_TICKER" not in state.market_data
        assert "LONG_TICKER" in state.market_data
        assert all(att.ticker == "LONG_TICKER" for att in state.attempts_log)

    def test_cli_synthetic_execution(self, tmp_path: Path):
        """Verify main() CLI runs with --synthetic and writes the Excel output."""
        out_file = tmp_path / "cli_out.xlsx"
        rc = main([
            "--synthetic",
            "--tickers", "SYN_A,SYN_B",
            "--max-attempts", "4",
            "--seed", "7",
            "--output", str(out_file),
        ])
        assert rc == 0
        assert out_file.exists()
        wb = load_workbook(out_file)
        assert set(wb.sheetnames) == {"Summary", "Risk Detail", "Portfolio"}

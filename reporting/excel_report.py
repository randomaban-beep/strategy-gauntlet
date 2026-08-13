"""
Excel report generator using openpyxl.

Creates a workbook with three sheets:
1. Summary: one row per candidate with approval status, rejection stage/reason, and headline metrics
2. Risk Detail: one row per candidate per simulated model (prob of profit, Sharpe, Sortino, VaR99, CVaR99)
3. Portfolio: HRP weights for approved strategies (or status message if < 2 approved)
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from pipeline.state import PipelineState


def write_excel_report(state: PipelineState, output_path: str | Path = "run_report.xlsx") -> Path:
    """Write the pipeline state to a formatted Excel workbook with 3 sheets."""
    path = Path(output_path)
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    # 1. Summary sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    summary_headers = [
        "Attempt #",
        "Ticker",
        "Strategy",
        "Params",
        "Approved",
        "Rejection Stage",
        "Rejection Reason",
        "IS Total Return",
        "IS Sharpe",
        "OOS Total Return",
        "OOS Sharpe",
        "Trusted Models",
        "Observed Sharpe (DSR)",
        "Expected Max Sharpe",
        "DSR Probability",
    ]
    ws_summary.append(summary_headers)
    for cell in ws_summary[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for idx, att in enumerate(state.attempts_log, start=1):
        is_ret = att.backtest_result.total_return if att.backtest_result else None
        is_sr = att.backtest_result.sharpe if att.backtest_result else None
        oos_ret = att.validation_result.total_return if att.validation_result else None
        oos_sr = att.validation_result.aggregate_sharpe if att.validation_result else None
        trusted = ", ".join(att.trusted_models.keys()) if att.trusted_models else ""
        obs_sr = att.scoring_result.observed_sharpe if att.scoring_result else None
        exp_max_sr = att.scoring_result.expected_max_sharpe if att.scoring_result else None
        dsr_prob = att.scoring_result.dsr_probability if att.scoring_result else None

        ws_summary.append([
            idx,
            att.ticker,
            att.strategy_name,
            json.dumps(att.params, sort_keys=True),
            "YES" if att.approved else "NO",
            att.rejection_stage or "",
            att.rejection_reason or "",
            round(is_ret, 4) if is_ret is not None else "",
            round(is_sr, 4) if is_sr is not None else "",
            round(oos_ret, 4) if oos_ret is not None else "",
            round(oos_sr, 4) if oos_sr is not None else "",
            trusted,
            round(obs_sr, 4) if obs_sr is not None else "",
            round(exp_max_sr, 4) if exp_max_sr is not None else "",
            round(dsr_prob, 4) if dsr_prob is not None else "",
        ])

    # 2. Risk Detail sheet
    ws_risk = wb.create_sheet(title="Risk Detail")
    risk_headers = [
        "Attempt #",
        "Ticker",
        "Strategy",
        "Model",
        "Prob of Profit",
        "Mean Total Return",
        "Sharpe",
        "Sortino",
        "Max Drawdown",
        "VaR99",
        "CVaR99",
    ]
    ws_risk.append(risk_headers)
    for cell in ws_risk[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for idx, att in enumerate(state.attempts_log, start=1):
        for model_name, metrics in att.risk_metrics.items():
            ws_risk.append([
                idx,
                att.ticker,
                att.strategy_name,
                model_name,
                round(metrics.prob_of_profit, 4),
                round(metrics.mean_total_return, 4),
                round(metrics.sharpe, 4),
                round(metrics.sortino, 4),
                round(metrics.max_drawdown, 4),
                round(metrics.var99, 4),
                round(metrics.cvar99, 4),
            ])

    # 3. Portfolio sheet
    ws_port = wb.create_sheet(title="Portfolio")
    port_headers = ["Strategy Key", "HRP Weight", "Portfolio Status"]
    ws_port.append(port_headers)
    for cell in ws_port[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    if state.portfolio_result and state.portfolio_result.weights:
        for key, weight in state.portfolio_result.weights.items():
            ws_port.append([key, round(weight, 6), state.portfolio_status])
    else:
        ws_port.append(["N/A", 0.0, state.portfolio_status])

    # Auto-size columns slightly for readability
    for ws in (ws_summary, ws_risk, ws_port):
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(path)
    return path
